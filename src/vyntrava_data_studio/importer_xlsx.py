"""Strict XLSX import into the shared textual table."""

from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path, PurePosixPath
import math
import stat
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from .importer import CsvTable


MAX_XLSX_BYTES = 10_485_760
MAX_NUMERIC_SIGNIFICANT_DIGITS = 15


class XlsxError(ValueError):
    def __init__(self, code: str, sheet: str | None = None,
                 row: int | None = None, column: int | None = None):
        self.code = code
        self.sheet = sheet
        self.row = row
        self.column = column
        location = "" if row is None else f" at row {row}, column {column}"
        super().__init__(code + location)


def _check_zip(path: Path) -> None:
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 1000:
                raise XlsxError("zip_limit")
            seen = set()
            total = 0
            for item in entries:
                name = item.filename
                parts = PurePosixPath(name).parts
                if (not name or name.startswith("/") or "\\" in name or ":" in name
                        or ".." in parts or name in seen):
                    raise XlsxError("invalid_zip_entry")
                seen.add(name)
                if item.file_size > 32 * 1024 * 1024:
                    raise XlsxError("zip_limit")
                total += item.file_size
                if total > 64 * 1024 * 1024:
                    raise XlsxError("zip_limit")
                if name.startswith("xl/externalLinks/"):
                    raise XlsxError("external_link")
                if name.lower().endswith("vbaproject.bin"):
                    raise XlsxError("macro_content")
    except (BadZipFile, OSError, ValueError) as exc:
        if isinstance(exc, XlsxError):
            raise
        raise XlsxError("invalid_workbook") from None


def _cell_text(cell, sheet: str) -> str:
    if cell.data_type in {"f", "e"} or cell.hyperlink is not None:
        raise XlsxError("invalid_cell", sheet, cell.row, cell.column)
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (date, datetime, time, timedelta)) or cell.is_date:
        raise XlsxError("invalid_cell", sheet, cell.row, cell.column)
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float, Decimal)):
        if isinstance(value, float) and not math.isfinite(value):
            raise XlsxError("invalid_cell", sheet, cell.row, cell.column)
        numeric = Decimal(str(value))
        if not numeric.is_finite():
            raise XlsxError("invalid_cell", sheet, cell.row, cell.column)
        digits = numeric.as_tuple().digits
        significant = len(digits)
        while significant > 1 and digits[significant - 1] == 0:
            significant -= 1
        if significant > MAX_NUMERIC_SIGNIFICANT_DIGITS:
            raise XlsxError("numeric_precision", sheet, cell.row, cell.column)
        return str(value)
    raise XlsxError("invalid_cell", sheet, cell.row, cell.column)


def read_xlsx(path: Path, sheet: str) -> CsvTable:
    """Read one visible sheet, rejecting features that change textual meaning."""
    try:
        source = Path(path)
        info = source.lstat()
        if source.suffix != ".xlsx" or not stat.S_ISREG(info.st_mode):
            raise XlsxError("invalid_file")
        if info.st_size > MAX_XLSX_BYTES:
            raise XlsxError("too_large")
    except (OSError, TypeError, ValueError) as exc:
        if isinstance(exc, XlsxError):
            raise
        raise XlsxError("invalid_file") from None
    _check_zip(source)
    try:
        book = load_workbook(source, read_only=False, data_only=False,
                             keep_links=False)
    except Exception:
        raise XlsxError("invalid_workbook") from None
    try:
        if sheet not in book.sheetnames or book[sheet].sheet_state != "visible":
            raise XlsxError("invalid_sheet")
        ws = book[sheet]
        if ws.max_row > 100_001 or ws.max_column > 256:
            raise XlsxError("table_limit")
        if ws.merged_cells.ranges:
            raise XlsxError("merged_cells")
        for index in range(1, ws.max_row + 1):
            if ws.row_dimensions[index].hidden:
                raise XlsxError("hidden_row", sheet, index, 1)
        for index in range(1, ws.max_column + 1):
            if ws.column_dimensions[ws.cell(1, index).column_letter].hidden:
                raise XlsxError("hidden_column", sheet, 1, index)
        records = []
        for cells in ws.iter_rows(min_row=1, max_row=ws.max_row,
                                  max_col=ws.max_column):
            if not records and any(cell.value is not None
                                   and not isinstance(cell.value, str) for cell in cells):
                raise XlsxError("invalid_header", sheet, 1, 1)
            row = tuple(_cell_text(cell, sheet) for cell in cells)
            if len(records) and not any(row):
                raise XlsxError("blank_record", sheet, len(records) + 1, 1)
            records.append(row)
        if not records or not records[0] or any(not h for h in records[0]) \
                or len(set(records[0])) != len(records[0]):
            raise XlsxError("invalid_header", sheet, 1, 1)
        return CsvTable(records[0], tuple(records[1:]), False, ",", "xlsx")
    finally:
        book.close()
