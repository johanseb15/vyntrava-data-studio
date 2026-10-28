import hashlib
import hmac
import json

import pytest
from openpyxl import Workbook, load_workbook

from vyntrava_data_studio.importer_xlsx import XlsxError, read_xlsx
from vyntrava_data_studio.recipes import RecipeError, load_recipe
from vyntrava_data_studio.sprint02 import run_recipe
from vyntrava_data_studio.package_verifier import verify_recipe_package


def workbook(path, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = "Ventas"
    for row in rows:
        sheet.append(row)
    book.save(path)
    book.close()


def test_reads_visible_sheet_as_text(tmp_path):
    path = tmp_path / "sales.xlsx"
    workbook(path, [("producto", "cantidad", "precio"),
                    ("Libro", 2, "10.00"), ("Lápiz", 3, "2.50")])
    table = read_xlsx(path, "Ventas")
    assert table.headers == ("producto", "cantidad", "precio")
    assert table.rows == (("Libro", "2", "10.00"), ("Lápiz", "3", "2.50"))


@pytest.mark.parametrize("value,observed", [
    (123456789012345.67, "123456789012345.7"),
    (0.12345678901234567, "0.1234567890123457"),
])
def test_rejects_numeric_cells_with_more_than_15_significant_digits(
        tmp_path, value, observed):
    path = tmp_path / "precision.xlsx"
    workbook(path, [("amount",), (value,)])
    book = load_workbook(path)
    try:
        assert str(book.active["A2"].value) == observed
    finally:
        book.close()
    with pytest.raises(XlsxError) as error:
        read_xlsx(path, "Ventas")
    assert (error.value.code, error.value.sheet, error.value.row,
            error.value.column) == ("numeric_precision", "Ventas", 2, 1)
    assert observed not in str(error.value)


@pytest.mark.parametrize("value,expected", [
    (0.10, "0.1"),
    (12.345, "12.345"),
    (1234567890123.45, "1234567890123.45"),
    (1e20, "1e+20"),
    ("123456789012345.67", "123456789012345.67"),
    ("0.12345678901234567", "0.12345678901234567"),
])
def test_preserves_numeric_values_within_limit_and_precise_text(
        tmp_path, value, expected):
    path = tmp_path / "precision.xlsx"
    workbook(path, [("amount",), (value,)])
    assert read_xlsx(path, "Ventas").rows == ((expected,),)


def test_rejects_missing_sheet_and_blank_middle_row(tmp_path):
    path = tmp_path / "sales.xlsx"
    workbook(path, [("name", "amount"), ("A", 2), (None, None), ("B", 3)])
    with pytest.raises(XlsxError) as error:
        read_xlsx(path, "Missing")
    assert error.value.code == "invalid_sheet"
    with pytest.raises(XlsxError) as error:
        read_xlsx(path, "Ventas")
    assert (error.value.code, error.value.row) == ("blank_record", 3)


@pytest.mark.parametrize("change,code", [
    (lambda sheet: setattr(sheet["A2"], "value", "=1+1"), "invalid_cell"),
    (lambda sheet: setattr(sheet["A2"], "hyperlink", "https://example.test"), "invalid_cell"),
    (lambda sheet: sheet.row_dimensions[2].__setattr__("hidden", True), "hidden_row"),
    (lambda sheet: sheet.merge_cells("A2:B2"), "merged_cells"),
    (lambda sheet: setattr(sheet, "sheet_state", "hidden"), "invalid_sheet"),
])
def test_rejects_ambiguous_workbook_features(tmp_path, change, code):
    path = tmp_path / "sales.xlsx"
    workbook(path, [("name", "amount"), ("A", 2)])
    book = __import__("openpyxl").load_workbook(path)
    if code == "invalid_sheet":
        book.create_sheet("Other")
    change(book.active)
    book.save(path)
    book.close()
    with pytest.raises(XlsxError) as error:
        read_xlsx(path, "Ventas")
    assert error.value.code == code


def test_recipe_xlsx_requires_minor_version_and_sheet(tmp_path):
    recipe = {"id": "sales", "version": "1.1.0", "tipo_encargo": "sales",
              "entradas": [{"id": "sales", "ruta": "sales.xlsx", "formato": "xlsx",
                            "hoja": "Ventas"}],
              "esquema": {"sales": {"columnas": ["name"], "extras": "rechazar"}},
              "parametros": {}, "reglas": [], "validaciones": [], "analisis": [],
              "salidas": [{"id": "summary", "tipo": "resumen_json", "ruta": "summary.json"}]}
    path = tmp_path / "recipe.json"
    path.write_text(json.dumps(recipe), encoding="utf-8")
    assert load_recipe(path).data["entradas"][0]["hoja"] == "Ventas"
    recipe["version"] = "1.0.0"
    path.write_text(json.dumps(recipe), encoding="utf-8")
    with pytest.raises(RecipeError) as error:
        load_recipe(path)
    assert error.value.code == "unsupported_version"


def test_xlsx_recipe_runs_analysis_and_records_sheet(tmp_path):
    root = tmp_path / "clients"
    job = root / "synthetic" / "sales"
    for folder in ("originales", "trabajo", "salidas", "evidencia"):
        (job / folder).mkdir(parents=True)
    source = job / "originales" / "sales.xlsx"
    workbook(source, [("product", "quantity", "price"),
                      ("Book", 2, "10.00"), ("Pencil", 3, "2.50")])
    original = source.read_bytes()
    recipe = {"id": "sales", "version": "1.1.0", "tipo_encargo": "sales",
              "entradas": [{"id": "sales", "ruta": "sales.xlsx", "formato": "xlsx",
                            "hoja": "Ventas"}],
              "esquema": {"sales": {"columnas": ["product", "quantity", "price"],
                                   "extras": "rechazar"}},
              "parametros": {}, "reglas": [], "validaciones": [],
              "analisis": [{"id": "total", "op": "suma_producto", "entrada": "sales",
                            "params": {"cantidad": "quantity", "precio": "price"}}],
              "salidas": [{"id": "summary", "tipo": "resumen_json", "ruta": "summary.json"}]}
    path = tmp_path / "recipe.json"
    path.write_text(json.dumps(recipe), encoding="utf-8")
    result = run_recipe(path, root, "synthetic", "sales")
    assert verify_recipe_package(path, root, "synthetic", "sales").ok
    report = json.loads(result.outputs["summary"].read_bytes())
    manifest = json.loads(result.manifest_path.read_bytes())
    assert report["analisis"]["total"]["total"] == "27.50"
    assert manifest["entradas"]["sales"]["hoja"] == "Ventas"
    assert source.read_bytes() == original


def test_pilot_b_xlsx_preserves_approval_analysis_png_and_manifest(tmp_path):
    root = tmp_path / "clients"
    job = root / "synthetic" / "sales"
    for folder in ("originales", "trabajo", "salidas", "evidencia"):
        (job / folder).mkdir(parents=True)
    source = job / "originales" / "sales.xlsx"
    workbook(source, [
        ("id", "fecha", "producto", "cantidad", "precio"),
        ("V001", "2026-01-05", "Libro", 2, "10.00"),
        ("V002", "2026-01-05", "Lápiz", 3, "2.50"),
        ("V003", "2026-01-06", "Libro", 1, "10.00"),
        ("V003", "2026-01-06", "Libro", 1, "10.00"),
        ("V004", "2026-01-06", "Libro", 0, "10.00"),
        ("V005", "2026-01-07", "Lápiz", 2, "2.50"),
    ])
    original = source.read_bytes()
    recipe_data = {
        "id": "pilot_b", "version": "1.1.0", "tipo_encargo": "sales",
        "entradas": [{"id": "sales", "ruta": "sales.xlsx", "formato": "xlsx",
                      "hoja": "Ventas"}],
        "esquema": {"sales": {"columnas": ["id", "fecha", "producto",
                                           "cantidad", "precio"],
                              "extras": "rechazar"}},
        "parametros": {},
        "reglas": [{"id": "deduplicate", "op": "eliminar_duplicados_exactos",
                    "entrada": "sales", "params": {"aprobacion": "approved"}}],
        "validaciones": [{"id": "positive_quantity", "op": "entero_positivo",
                          "entrada": "sales", "params": {"columna": "cantidad"},
                          "severidad": "advertir"}],
        "analisis": [{"id": "by_date", "op": "suma_producto", "entrada": "sales",
                      "params": {"cantidad": "cantidad", "precio": "precio",
                                 "grupo": "fecha"}}],
        "salidas": [{"id": "summary", "tipo": "resumen_json", "ruta": "summary.json"},
                    {"id": "chart", "tipo": "grafico_png", "ruta": "sales.png",
                     "fuente": "by_date"}],
    }
    recipe = tmp_path / "recipe.json"
    recipe.write_text(json.dumps(recipe_data, ensure_ascii=False), encoding="utf-8")
    key = b"synthetic-test-key"
    approval = {
        "id": "approved", "responsable": "synthetic-reviewer",
        "receta_sha256": hashlib.sha256(recipe.read_bytes()).hexdigest(),
        "entradas_sha256": {"sales": hashlib.sha256(original).hexdigest()},
    }
    payload = json.dumps(approval, sort_keys=True, separators=(",", ":"),
                         ensure_ascii=False).encode("utf-8")
    approval["firma"] = hmac.new(key, payload, hashlib.sha256).hexdigest()

    result = run_recipe(recipe, root, "synthetic", "sales",
                        approvals={"approved": approval}, approval_key=key)
    summary = json.loads(result.outputs["summary"].read_bytes())
    analysis = summary["analisis"]["by_date"]
    chart = result.outputs["chart"].read_bytes()
    manifest = json.loads(result.manifest_path.read_bytes())

    assert result.status == "revision_requerida"
    assert (analysis["filas_incluidas"], analysis["filas_excluidas"],
            analysis["unidades"], analysis["total"]) == (4, 1, 8, "42.50")
    assert analysis["grupos"] == {"2026-01-05": "27.50",
                                   "2026-01-06": "10.00",
                                   "2026-01-07": "5.00"}
    assert chart.startswith(b"\x89PNG\r\n\x1a\n")
    assert manifest["entradas"]["sales"] == {
        "sha256": hashlib.sha256(original).hexdigest(),
        "formato": "xlsx", "hoja": "Ventas"}
    assert manifest["salidas"]["chart"]["sha256"] == hashlib.sha256(chart).hexdigest()
    assert manifest["aprobaciones"]["approved"]["firma"] == approval["firma"]
    assert source.read_bytes() == original
